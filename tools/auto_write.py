# tools/auto_write.py
# -*- coding: utf-8 -*-
r"""
재다몰 자동 업로드 (로그인 보장형 단일 스크립트) - SmartEditor2(memo) 대응

✅ 현재 게시판(boardid=41) 기준 확인된 구조
- oEditors.getById['memo'] 존재
- textarea id/name = memo
- iframe id/name 비어있을 수 있음(= se2_iframe 같은 고정 id 없음)

동작 요약
1) 시크릿 크롬 실행
2) 로그인 페이지로 이동 → #login_id/#login_pw 자동 입력 → 엔터 submit
3) 글쓰기 페이지 진입 (write.php?boardid=41)
4) 제목 입력
5) 본문 입력 (HTML)
   - 1순위: oEditors.getById['memo'] → SET_CONTENTS + UPDATE_CONTENTS_FIELD
   - 2순위: textarea#memo value 세팅(제출값 직접)
6) 등록 버튼 클릭(JS click)
7) 브라우저는 기본적으로 열어둠(옵션 --close로 닫기)

.env (프로젝트 루트에 위치)
ZAEDA_ID=...
ZAEDA_PW=...

실행 예시
- 테스트(엑셀 없이):
  python tools/auto_write.py --no-excel

- 엑셀 있으면(없으면 자동으로 no-excel처럼 동작):
  python tools/auto_write.py
"""

from dotenv import load_dotenv
import os
import argparse
import datetime
import time
import traceback
from pathlib import Path

import openpyxl
from selenium.webdriver import Chrome, ChromeOptions
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
from selenium.common.exceptions import TimeoutException, UnexpectedAlertPresentException


# ──────────────────────────────
# Constants / Paths
# ──────────────────────────────
ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
XLSX = DOCS / "data.xlsx"

MAX_WAIT = 20

LOGIN_URL = "https://zae-da.com/bbs/login.php?url=%2F"
DEFAULT_LIST_URL = "https://zae-da.com/bbs/list.php?boardid=41"
DEFAULT_WRITE_URL = "https://zae-da.com/bbs/write.php?boardid=41"


# ──────────────────────────────
# Utils
# ──────────────────────────────
def log(msg: str):
    print(msg, flush=True)


def wait_ready(drv, timeout=MAX_WAIT):
    WebDriverWait(drv, timeout).until(
        lambda d: d.execute_script("return document.readyState") == "complete"
    )


def accept_all_alerts(drv, limit=3):
    for _ in range(limit):
        try:
            a = drv.switch_to.alert
            txt = a.text
            log(f"⚠ 알럿 감지: {txt}")
            a.accept()
            time.sleep(0.3)
        except Exception:
            break


# ──────────────────────────────
# Excel (optional)
# ──────────────────────────────
def load_next_row():
    if not XLSX.exists():
        raise FileNotFoundError(f"엑셀 파일이 없습니다: {XLSX}")

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active

    for i in range(2, ws.max_row + 1):
        title = (ws[f"A{i}"].value or "").strip()
        body = (ws[f"B{i}"].value or "").strip()
        status = (ws[f"C{i}"].value or "").strip().upper()
        if title and body and status not in ("DONE", "PUBLISHED", "SKIP"):
            return wb, ws, i, title, body

    return wb, ws, None, None, None


def mark_done(wb, ws, row: int):
    ws[f"C{row}"] = "DONE"
    ws[f"D{row}"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    wb.save(XLSX)


# ──────────────────────────────
# Driver
# ──────────────────────────────
def build_options():
    opts = ChromeOptions()
    opts.add_argument("--incognito")
    opts.add_argument("--start-maximized")
    opts.add_argument("--no-first-run")
    opts.add_argument("--no-default-browser-check")
    opts.add_argument("--disable-extensions")
    opts.add_argument("--disable-popup-blocking")
    opts.add_argument("--disable-dev-shm-usage")
    opts.add_argument("--disable-gpu")
    opts.add_argument("--remote-allow-origins=*")
    opts.add_experimental_option("excludeSwitches", ["enable-automation"])
    opts.add_experimental_option("useAutomationExtension", False)
    return opts


def setup_driver():
    opts = build_options()
    drv = Chrome(service=Service(ChromeDriverManager().install()), options=opts)
    return drv, "incognito"


# ──────────────────────────────
# Login
# ──────────────────────────────
def is_logged_in(drv) -> bool:
    try:
        logout = drv.find_elements(
            By.XPATH,
            "//a[contains(.,'로그아웃') or contains(.,'Logout') or contains(@href,'logout')]",
        )
        return bool(logout)
    except Exception:
        return False


def try_auto_login(drv) -> bool:
    uid = os.getenv("ZAEDA_ID", "").strip()
    pw = os.getenv("ZAEDA_PW", "").strip()

    if not uid or not pw:
        log("[login] ZAEDA_ID/ZAEDA_PW 미설정 → 자동 로그인 스킵")
        return False

    drv.get(LOGIN_URL)
    wait_ready(drv, 15)
    accept_all_alerts(drv)

    id_el = WebDriverWait(drv, 10).until(
        EC.presence_of_element_located((By.CSS_SELECTOR, "#login_id"))
    )
    pw_el = WebDriverWait(drv, 10).until(
        EC.presence_of_element_located((By.CSS_SELECTOR, "#login_pw"))
    )

    id_el.clear()
    id_el.send_keys(uid)
    pw_el.clear()
    pw_el.send_keys(pw)
    pw_el.send_keys("\n")

    time.sleep(1.0)
    accept_all_alerts(drv)
    time.sleep(0.5)

    if is_logged_in(drv):
        log("🔐 자동 로그인 성공")
        return True

    return False


def ensure_login(drv, list_url_for_check: str):
    try:
        drv.get(list_url_for_check)
        wait_ready(drv, 15)
        accept_all_alerts(drv)
    except Exception:
        pass

    if is_logged_in(drv):
        log("🔓 이미 로그인 상태입니다.")
        return

    if try_auto_login(drv):
        return

    log("👤 자동 로그인 실패 → 수동 로그인 안내")
    drv.get(LOGIN_URL)
    wait_ready(drv, 15)
    accept_all_alerts(drv)

    input("재다몰에 수동 로그인 후 엔터를 눌러 계속하세요... ")

    drv.get(list_url_for_check)
    wait_ready(drv, 15)
    accept_all_alerts(drv)

    if not is_logged_in(drv):
        raise RuntimeError("로그인이 확인되지 않았습니다.")


# ──────────────────────────────
# Navigation (write page)
# ──────────────────────────────
def find_subject(drv):
    for sel in [
        "input[name='wr_subject']",
        "input#wr_subject",
        "input[name='subject']",
        "input#subject",
        "input[name='title']",
        "input#title",
    ]:
        els = drv.find_elements(By.CSS_SELECTOR, sel)
        if els:
            return els[0]
    return None


def ensure_write_page(drv, list_url: str, write_url: str) -> None:
    # 1) write_url 직접 접근
    drv.get(write_url)
    wait_ready(drv, 15)
    accept_all_alerts(drv)

    if find_subject(drv) is not None:
        log("✅ 글쓰기 페이지(직접 URL) 진입 성공")
        return

    # 2) 목록에서 글쓰기 버튼 클릭(보험)
    drv.get(list_url)
    wait_ready(drv, 15)
    accept_all_alerts(drv)

    btns = drv.find_elements(By.CSS_SELECTOR, "a.btn_lsmall, a[href*='write.php']")
    for b in btns:
        try:
            if b.is_enabled() and ("글쓰기" in (b.text or "")):
                b.click()
                time.sleep(0.5)
                accept_all_alerts(drv)
                wait_ready(drv, 15)
                if find_subject(drv) is not None:
                    log("✅ 글쓰기 페이지(리스트→버튼) 진입 성공")
                    return
        except Exception:
            continue

    raise RuntimeError("글쓰기 페이지로 진입하지 못했습니다.")


# ──────────────────────────────
# Content builder (HTML) - 줄간격/굵기/이미지
# ──────────────────────────────
def build_test_post():
    """
    ✅ 줄간격(line-height) 고정 + 제목/소제목 굵게 + 이미지 2장 포함(테스트용)
    - 이모지는 환경에 따라 ????로 깨질 수 있어서, 테스트는 기호(★/✓/>) 위주로 사용
    """
    title = "다이어트/비만 관리 | 운동해도 변화 없을 때, 먼저 점검할 3가지"

    img1 = "https://images.unsplash.com/photo-1554284126-aa88f22d8b74?auto=format&fit=crop&w=1200&q=80"
    img2 = "https://images.unsplash.com/photo-1543362906-acfc16c67564?auto=format&fit=crop&w=1200&q=80"

    # 핵심: wrapper에 font-size/line-height 지정 (행간 줄이기)
    body_html = f"""
<div style="font-size:15px; line-height:1.45; color:#111;">
  <p style="margin:0 0 10px 0;">
    분명 예전과 비슷하게 먹고, 운동도 꾸준히 하는데 체중이 잘 안 내려갈 때가 있죠.<br/>
    그럴 땐 “의지가 부족해서”라기보다, 몸이 현재 어떤 신호를 보내고 있는지 먼저 확인해보는 게 도움이 됩니다.<br/>
    오늘은 생활 속에서 바로 점검할 수 있는 3가지를 정리해볼게요.
  </p>

  <p style="margin:12px 0 8px 0;">
    <img src="{img1}" alt="운동 이미지" style="max-width:100%; border-radius:10px;" />
  </p>

  <p style="margin:14px 0 6px 0;">
    <strong style="font-size:17px;">1) 수면 리듬부터 정리하기 ★</strong>
  </p>
  <p style="margin:0 0 10px 0;">
    잠이 부족하거나 취침 시간이 들쑥날쑥하면 식욕·포만감에 관여하는 균형이 흔들릴 수 있어요.<br/>
    특히 늦은 밤에 단 음식이 당기거나, 아침에 몸이 무겁게 느껴진다면 수면부터 조정해보는 게 좋습니다.<br/>
    오늘부터는 <strong>기상 시간</strong>을 고정하고, 취침 1시간 전에는 화면 밝기(휴대폰/노트북)를 낮춰보세요.<br/>
    <strong>초보자 팁:</strong> 갑자기 2시간 당기기보다 15~30분씩 앞당기면 훨씬 덜 힘들어요.
  </p>

  <p style="margin:14px 0 6px 0;">
    <strong style="font-size:17px;">2) ‘가공당 + 음료’부터 줄이기 ✓</strong>
  </p>
  <p style="margin:0 0 10px 0;">
    같은 칼로리라도 액상 형태(달달한 커피, 주스, 탄산)는 포만감이 낮아서 과식으로 이어지기 쉬워요.<br/>
    또 자주 마시면 몸 컨디션이 들쑥날쑥해질 수 있습니다.<br/>
    <strong>실행법:</strong> 일주일만 음료를 물/무가당 차로 바꿔보고, 디저트는 주 2회로 요일을 정해보세요.<br/>
    <strong>초보자 팁:</strong> 완전 금지 대신 ‘요일 제한’이 오래 갑니다.
  </p>

  <p style="margin:12px 0 8px 0;">
    <img src="{img2}" alt="식단 이미지" style="max-width:100%; border-radius:10px;" />
  </p>

  <p style="margin:14px 0 6px 0;">
    <strong style="font-size:17px;">3) 단백질·식이섬유를 ‘끼니마다’ 넣기 &gt;</strong>
  </p>
  <p style="margin:0 0 10px 0;">
    식사에서 단백질과 식이섬유가 부족하면 금방 배가 고파지고, 간식으로 이어지기 쉬워요.<br/>
    특히 아침을 빵/커피로 끝내는 날이 많다면 점심 전에 허기가 크게 올 수 있습니다.<br/>
    <strong>실행법:</strong> 끼니마다 단백질(달걀, 두부, 생선, 살코기)과 채소(나물/샐러드/김치)를 한 가지씩만 추가해보세요.<br/>
    <strong>초보자 팁:</strong> ‘계란 1개 + 채소 반찬 1개’만 고정해도 충분해요.
  </p>

  <p style="margin:14px 0 6px 0;"><strong style="font-size:16px;">주의사항</strong></p>
  <p style="margin:0 0 10px 0;">
    기저질환(당뇨, 갑상선 질환 등)이 있거나 약물을 복용 중이라면 체중 변화가 다르게 나타날 수 있어요.<br/>
    또 너무 적게 먹고 많이 운동하면 오히려 피로가 쌓여 지속이 어려울 수 있으니 ‘지속 가능한 범위’로 조절해보세요.
  </p>

  <p style="margin:14px 0 6px 0;"><strong style="font-size:16px;">요약</strong></p>
  <p style="margin:0 0 10px 0;">
    운동만 늘리기 전에 수면 리듬·음료/가공당·끼니 구성(단백질/식이섬유) 3가지를 먼저 점검해보면 좋습니다.<br/>
    작은 습관 하나가 식욕과 컨디션을 바꾸고, 장기적으로 체중 관리에도 도움이 될 수 있어요.
  </p>

  <p style="margin:14px 0 6px 0;"><strong style="font-size:16px;">근거자료(참고)</strong></p>
  <p style="margin:0;">
    - WHO: 식생활 및 만성질환 관련 자료<br/>
    - 질병관리청: 비만/대사건강 관련 건강정보<br/>
    - 대한비만학회: 비만 진료지침(생활습관 관리)
  </p>

  <hr style="border:none;border-top:1px solid #ddd; margin:14px 0;" />

  <p style="margin:0; color:#444;">
    이 글은 일반적인 건강 정보를 제공하기 위한 것이며, 의료적 진단이나 치료를 대신하지 않습니다. 개인별 상태에 따라 전문가 상담이 필요할 수 있습니다.
  </p>
</div>
""".strip()

    return title, body_html


# ──────────────────────────────
# SmartEditor2(memo) body set
# ──────────────────────────────
def wait_editor_ready(drv):
    WebDriverWait(drv, 20).until(
        lambda d: d.find_elements(By.CSS_SELECTOR, "textarea#memo, textarea[name='memo']")
        or d.execute_script(
            "return (typeof window.oEditors !== 'undefined') && window.oEditors && oEditors.getById && oEditors.getById['memo'];"
        )
    )


def set_body_oeditors_memo(drv, body_html: str) -> str:
    return drv.execute_script(
        """
        const html = arguments[0];
        try {
            if (window.oEditors && oEditors.getById && oEditors.getById['memo']) {
                oEditors.getById['memo'].exec('SET_CONTENTS', [html]);
                oEditors.getById['memo'].exec('UPDATE_CONTENTS_FIELD', []);
                return 'OK:memo';
            }
            return 'NO_EDITORS';
        } catch(e) {
            return 'ERR:' + e.toString();
        }
        """,
        body_html,
    )


def set_body_textarea_memo(drv, body_html: str) -> bool:
    tas = drv.find_elements(By.CSS_SELECTOR, "textarea#memo, textarea[name='memo']")
    if not tas:
        return False
    drv.execute_script("arguments[0].value = arguments[1];", tas[0], body_html)
    return True


def sync_editor_before_submit(drv):
    try:
        drv.execute_script(
            """
            try {
                if (window.oEditors && oEditors.getById && oEditors.getById['memo']) {
                    oEditors.getById['memo'].exec('UPDATE_CONTENTS_FIELD', []);
                }
            } catch(e) {}
            """
        )
    except Exception:
        pass


def set_post_body(drv, body_html: str):
    log("본문 입력 시작… (memo)")
    wait_editor_ready(drv)

    # 1) oEditors 먼저(화면/textarea 동기화까지 되는 편)
    res = "NO_EDITORS"
    try:
        res = set_body_oeditors_memo(drv, body_html)
        log(f"oEditors 결과: {res}")
    except Exception as e:
        log(f"⚠ oEditors 예외(무시): {repr(e)}")

    # 2) textarea도 한 번 더(제출값 보장)
    if set_body_textarea_memo(drv, body_html):
        log("본문 입력 완료 ✓ (textarea#memo)")

    # 최종 검증
    v = drv.execute_script(
        "const t=document.querySelector('textarea#memo,textarea[name=memo]'); return t ? (t.value || '') : null;"
    )
    if v is None or len(v.strip()) == 0:
        raise RuntimeError("본문 입력 검증 실패: textarea#memo 값이 비어있습니다.")


# ──────────────────────────────
# Main
# ──────────────────────────────
def main():
    load_dotenv()

    ap = argparse.ArgumentParser()
    ap.add_argument("--url", default=DEFAULT_WRITE_URL, help="글쓰기 URL")
    ap.add_argument("--list-url", default=DEFAULT_LIST_URL, help="글 목록 URL")
    ap.add_argument("--no-excel", action="store_true", help="엑셀 없이 테스트 본문/제목 사용")
    ap.add_argument("--close", action="store_true", help="작업 후 브라우저 닫기")
    args = ap.parse_args()

    drv, profile_info = setup_driver()
    log(f"프로필 정보: {profile_info}")

    try:
        ensure_login(drv, args.list_url)
        ensure_write_page(drv, args.list_url, args.url)

        # ✅ 기본값 먼저 세팅(= UnboundLocalError 방지)
        wb = ws = row = None
        title, body_html = build_test_post()

        # 엑셀 모드 (단, 파일 없으면 자동으로 테스트로 진행)
        if not args.no_excel:
            try:
                wb, ws, row, title_from_xlsx, body_from_xlsx = load_next_row()
                if row:
                    title = title_from_xlsx
                    # 엑셀은 보통 텍스트일 테니, 최소 HTML wrapper로 감싸서 행간 고정
                    body_html = f"<div style='font-size:15px; line-height:1.45;'>{(body_from_xlsx or '').replace('\\n','<br/>')}</div>"
                else:
                    log("대기 중인 업로드 건이 없습니다. → 테스트 글로 진행합니다.")
            except FileNotFoundError as e:
                log(f"⚠ {e} → 테스트 글로 진행합니다.")

        # 제목 입력
        ti = find_subject(drv)
        if not ti:
            raise RuntimeError("제목 입력 필드를 찾을 수 없습니다.")
        ti.clear()
        ti.send_keys(title)
        log("제목 입력 완료 ✓")

        # 본문 입력
        set_post_body(drv, body_html)

        # 제출 버튼 찾기
        submit_btn = None
        for xp in [
            "//button[contains(.,'등록') or contains(.,'작성') or contains(.,'저장')]",
            "//input[@type='submit']",
        ]:
            btns = drv.find_elements(By.XPATH, xp)
            if btns:
                submit_btn = btns[0]
                break

        if submit_btn is None:
            raise RuntimeError("제출 버튼을 찾을 수 없습니다.")

        sync_editor_before_submit(drv)

        # JS 클릭
        drv.execute_script("arguments[0].click();", submit_btn)
        time.sleep(0.7)
        accept_all_alerts(drv)
        wait_ready(drv, 15)

        log("등록 버튼 클릭 ✓")

        # 엑셀 DONE 처리
        if wb and ws and row:
            mark_done(wb, ws, row)
            log("✅ 업로드 완료 → DONE 처리")

    except UnexpectedAlertPresentException:
        try:
            a = drv.switch_to.alert
            log(f"❌ 알럿으로 인해 중단되었습니다. 메시지: {a.text}")
            a.accept()
        except Exception:
            pass

    except Exception as e:
        log(f"❌ 오류 타입: {type(e)}")
        log(f"❌ 오류 repr: {repr(e)}")
        log("❌ 스택트레이스:")
        log(traceback.format_exc())

    finally:
        if args.close:
            try:
                drv.quit()
            except Exception:
                pass
        else:
            log("브라우저는 열어둡니다. 작업 내용을 확인하세요.")


if __name__ == "__main__":
    main()
